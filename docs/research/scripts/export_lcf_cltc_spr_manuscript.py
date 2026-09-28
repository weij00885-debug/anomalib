# /// script
# requires-python = ">=3.11"
# dependencies = ["pypandoc-binary>=1.15,<2", "python-docx>=1.1,<2"]
# ///
"""Export the research manuscript and data appendix as one editable Word file."""

from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
import pypandoc


def main() -> None:
    """Convert Markdown math, figures and data tables to an editable manuscript."""
    research = Path(__file__).resolve().parents[1]
    main_text = (research / "dinomaly_lcf_cltc_spr_paper.md").read_text(encoding="utf-8")
    data_text = (research / "dinomaly_lcf_cltc_spr_paper_data.md").read_text(encoding="utf-8")
    output = research / "dinomaly_lcf_cltc_spr_manuscript.docx"
    pypandoc.convert_text(
        main_text + "\n\n" + data_text,
        "docx",
        format="markdown+tex_math_single_backslash",
        outputfile=str(output),
        extra_args=[f"--resource-path={research}", "--standalone", "--toc", "--toc-depth=2",
                    "--metadata=title:LCF＋CLTC＋SPR 单类人脸防伪研究稿"],
    )
    doc = Document(output)
    for section in doc.sections:
        section.page_width, section.page_height = Cm(21), Cm(29.7)
        section.top_margin = section.bottom_margin = Cm(2)
        section.left_margin = section.right_margin = Cm(1.7)
    for style_name in ("Normal", "Body Text", "First Paragraph"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)
        style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "宋体")
    for paragraph in doc.paragraphs:
        if paragraph.text == "LCF＋CLTC＋SPR 论文数据附录":
            paragraph.paragraph_format.page_break_before = True
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    doc.settings.element.append(update_fields)
    for table in doc.tables:
        table.autofit = True
        if table.rows:
            header = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(header)
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8)
    doc.save(output)
    print(f"Saved {output}; tables={len(doc.tables)}, embedded figures={len(doc.inline_shapes)}")


if __name__ == "__main__":
    main()
